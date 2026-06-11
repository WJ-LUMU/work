# src/excel_exporter.py
"""
Excel导出模块

负责：
1. 创建新的Excel文件
2. 格式化工作表
3. 添加核对结果
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from typing import Dict
from datetime import datetime


class ExcelExporter:
    """Excel导出类"""
    
    def __init__(self):
        """初始化导出器"""
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def export(self, result: Dict, save_dir: Path) -> Path:
        """
        导出处理结果到Excel文件
        
        参数：
            result: 处理结果字典
            save_dir: 保存目录
        
        返回：
            输出文件的Path对象
        """
        # 创建工作簿
        wb = Workbook()
        wb.remove(wb.active)  # 删除默认工作表
        
        # 创建三个工作表
        ws1 = wb.create_sheet("按标准顺序整理", 0)
        ws2 = wb.create_sheet("新增80码", 1)
        ws3 = wb.create_sheet("核对结果", 2)
        
        # 填充各工作表
        self._fill_standard_sheet(ws1, result['standard_result'])
        self._fill_new_codes_sheet(ws2, result['new_codes'])
        self._fill_summary_sheet(ws3, result['summary'])
        
        # 生成文件名
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"80码整理结果_{timestamp}.xlsx"
        output_path = save_dir / filename
        
        # 保存文件
        wb.save(output_path)
        
        return output_path
    
    def _fill_standard_sheet(self, ws, data):
        """
        填充"按标准顺序整理"工作表
        
        参数：
            ws: 工作表对象
            data: 数据列表
        """
        # 设置列标题
        headers = ['序号', '80码', '合计完成', '汇总表是否存在']
        ws.append(headers)
        
        # 格式化标题行
        self._format_header_row(ws, len(headers))
        
        # 添加数据
        for item in data:
            ws.append([
                item['sequence'],
                item['code'],
                item['quantity'],
                item['exists']
            ])
        
        # 格式化数据行
        self._format_data_rows(ws, len(data), len(headers))
        
        # 调整列宽
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 16
        
        # 设置80码列为文本格式
        for row in range(2, len(data) + 2):
            cell = ws[f'B{row}']
            cell.number_format = '@'  # 文本格式
        
        # 设置数量列为整数格式
        for row in range(2, len(data) + 2):
            cell = ws[f'C{row}']
            cell.number_format = '0'
        
        # 冻结首行
        ws.freeze_panes = 'A2'
        
        # 启用自动筛选
        ws.auto_filter.ref = f'A1:D{len(data) + 1}'
    
    def _fill_new_codes_sheet(self, ws, data):
        """
        填充"新增80码"工作表
        
        参数：
            ws: 工作表对象
            data: 数据列表
        """
        if not data:
            # 如果没有新增80码，显示提示信息
            ws.append(['未发现新增80码'])
            ws['A1'].font = Font(size=12, bold=True, color="008000")
            ws.column_dimensions['A'].width = 20
        else:
            # 设置列标题
            headers = ['序号', '80码', '合计完成', '备注']
            ws.append(headers)
            
            # 格式化标题行
            self._format_header_row(ws, len(headers))
            
            # 添加数据
            for item in data:
                ws.append([
                    item['sequence'],
                    item['code'],
                    item['quantity'],
                    item['remark']
                ])
            
            # 格式化数据行
            self._format_data_rows(ws, len(data), len(headers))
            
            # 调整列宽
            ws.column_dimensions['A'].width = 8
            ws.column_dimensions['B'].width = 15
            ws.column_dimensions['C'].width = 12
            ws.column_dimensions['D'].width = 25
            
            # 设置80码列为文本格式
            for row in range(2, len(data) + 2):
                cell = ws[f'B{row}']
                cell.number_format = '@'
            
            # 设置数量列为整数格式
            for row in range(2, len(data) + 2):
                cell = ws[f'C{row}']
                cell.number_format = '0'
            
            # 冻结首行
            ws.freeze_panes = 'A2'
            
            # 启用自动筛选
            ws.auto_filter.ref = f'A1:D{len(data) + 1}'
    
    def _fill_summary_sheet(self, ws, summary):
        """
        填充"核对结果"工作表
        
        参数：
            ws: 工作表对象
            summary: 汇总数据字典
        """
        # 准备数据
        items = [
            ('项目', '数量'),
            ('原始汇总表总计', summary['original_total']),
            ('标准清单内总计', summary['standard_total']),
            ('新增80码总计', summary['new_total']),
            ('整理后总计', summary['final_total']),
            ('', ''),  # 空行
            ('数量差异', summary['difference']),
            ('核对结果', summary['match_status']),
        ]
        
        # 添加数据
        for idx, (label, value) in enumerate(items, 1):
            ws[f'A{idx}'] = label
            ws[f'B{idx}'] = value
        
        # 格式化
        # 标题行
        for col in ['A1', 'B1']:
            cell = ws[col]
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.thin_border
        
        # 数据行
        for row in range(2, len(items) + 1):
            for col in ['A', 'B']:
                cell = ws[f'{col}{row}']
                
                # 空行
                if row == 6:
                    continue
                
                # 标签列
                if col == 'A':
                    cell.font = Font(bold=True, size=10)
                    cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
                
                # 数值列
                else:
                    cell.font = Font(size=10)
                    
                    # 如果是数值，右对齐
                    if isinstance(cell.value, (int, float)):
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                        cell.number_format = '0'
                    
                    # 如果是"核对结果"，加色彩标识
                    if row == len(items):  # 最后一行
                        if summary['match_status'] == '核对一致':
                            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                            cell.font = Font(color="006100", size=10, bold=True)
                        else:
                            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                            cell.font = Font(color="9C0006", size=10, bold=True)
                
                cell.border = self.thin_border
        
        # 调整列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
    
    def _format_header_row(self, ws, col_count):
        """
        格式化表头行
        
        参数：
            ws: 工作表对象
            col_count: 列数
        """
        for col in range(1, col_count + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = self.thin_border
    
    def _format_data_rows(self, ws, row_count, col_count):
        """
        格式化数据行
        
        参数：
            ws: 工作表对象
            row_count: 行数
            col_count: 列数
        """
        for row in range(2, row_count + 2):
            for col in range(1, col_count + 1):
                cell = ws.cell(row=row, column=col)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = self.thin_border
                
                # 条纹行效果
                if row % 2 == 0:
                    cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
